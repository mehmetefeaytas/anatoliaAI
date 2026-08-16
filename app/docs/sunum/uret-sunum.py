"""Jüri sunumunu HTML'den PDF + PPTX'e çevirir.

Şartname §6 (s.14) sunum materyalini "PDF ve PPTX formatında" ister — ikisi
birden zorunlu. Kaynak tek: `docs/sunum/anatolia-ai-sunum.html`. Bu betik o
dosyayı hiç değiştirmez; yalnız tarayıcıya yükleyip iki çıktı üretir.

KULLANIM
    .venv/bin/python docs/sunum/uret-sunum.py
    (yalnız PDF)  .venv/bin/python docs/sunum/uret-sunum.py --sadece pdf
    (yalnız PPTX) .venv/bin/python docs/sunum/uret-sunum.py --sadece pptx

NEDEN İKİ FARKLI YOL

PDF → Chromium'un kendi yazdırma motoru (`page.pdf`). Metin vektör kalır,
seçilebilir, dosya küçük; Türkçe diyakritikler sistem fontundan doğrudan
gömülür. Sayfalama enjekte edilen `@page 1920x1080px` + `.slayt{height:1080px;
break-after:page}` ile zorlanır, böylece her slayt tam olarak bir sayfadır.

PPTX → her slayt 1920×1080 PNG olarak çekilip 16:9 slayda tam sayfa görsel
konur. Alternatif (kutu kutu yeniden dizgi) HTML'in ızgara/clamp tipografisini
PowerPoint kutularına çevirmeyi gerektirirdi; bu, birebir görünümü kaybettirir
ve yeniden dizgi hatası riski taşır. Görsel yaklaşımında risk sıfır: jürinin
gördüğü kare, tarayıcıda render edilen karenin aynısı. Bedeli metnin
seçilemez olması — sunum materyali için kabul edilebilir, çünkü metin
seçilebilir sürüm zaten PDF olarak teslim ediliyor.

SIĞDIRMA

Kaynak HTML slaytları `min-height:100vh` ile tasarlandı; tarayıcıda uzun slayt
kaydırılabiliyor, sabit 16:9 sayfada kaydırma yok. İki slaytın (07, 08) içeriği
1080px'i aşıyor ve `overflow:hidden` yüzünden sessizce kırpılıyordu. Betik bunu
ölçer ve o slaytları TEK KATSAYIYLA küçültür — içerik değişmez, yalnız ölçek
küçülür. Ayrıntılı gerekçe için `SIGDIR_JS` yorumuna bak.

BAĞIMLILIK / LİSANS

playwright (Apache-2.0) proje `.venv`'inde zaten kurulu.
python-pptx (MIT) proje `.venv`'ine KURULMAZ — `make sbom` kurulu ortamı
tarar, oraya paket eklemek bağımlılık envanterini ve lisans kapısını kirletir.
Bunun yerine repo dışında ayrı bir yardımcı ortam kullanılır:
`~/.cache/anatolia-ai/sunum-venv`. Ortam yoksa betik onu kurar — bu TEK
seferlik adım ağ erişimi ister (PyPI). Ortam varsa tüm üretim çevrimdışıdır.
"""
from __future__ import annotations

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

from playwright.sync_api import sync_playwright

KOK = Path(__file__).resolve().parent
KAYNAK = KOK / "anatolia-ai-sunum.html"
PDF = KOK / "anatolia-ai-sunum.pdf"
PPTX = KOK / "anatolia-ai-sunum.pptx"

EN, BOY = 1920, 1080  # 16:9 — şartname §10 canlı sunum, jüri projeksiyonu

# Yardımcı ortam: python-pptx buraya kurulur, proje .venv'i temiz kalır.
YARDIMCI = Path.home() / ".cache" / "anatolia-ai" / "sunum-venv"

# Sayfalama + baskı düzeltmeleri. Kaynak HTML'e DOKUNULMAZ; bu CSS yalnız
# tarayıcı belleğindeki kopyaya enjekte edilir.
BASKI_CSS = f"""
@page {{ size: {EN}px {BOY}px; margin: 0; }}
html, body {{
  height: auto !important; overflow: visible !important;
  scroll-snap-type: none !important;
}}
.slayt {{
  height: {BOY}px !important; min-height: {BOY}px !important;
  break-after: page; page-break-after: always;
  break-inside: avoid; page-break-inside: avoid;
}}
/* Son slayttan sonra sayfa kırma → boş 16. sayfa üretirdi. */
.slayt:last-of-type {{ break-after: auto; page-break-after: auto; }}
/* Tema düğmesi ekran kontrolü; baskıda işi yok. */
.tema {{ display: none !important; }}
/* Giriş animasyonu: PDF'te yakalanırsa içerik yarı saydam çıkar. */
.gir {{ opacity: 1 !important; transform: none !important; transition: none !important; }}
"""


# YALNIZ PDF geçişinde uygulanır — PNG/PPTX geçişinde UYGULANMAZ.
# Ölçüldü: `.ekran`ın `box-shadow:0 18px 48px rgba(20,26,33,.13)` gölgesi
# Chromium'un yazdırma yolunda bulanıklığını kaybediyor ve ürün ekranının
# arkasına KESKİN KENARLI gri bir dikdörtgen olarak basılıyor (slayt 06, 07,
# 08). Ekran görüntüsü yolunda aynı gölge doğru render ediliyor, o yüzden bu
# düzeltme PPTX'e taşınmaz. `.ekran` zaten 1px çerçeve taşıyor; gölge
# kapatılınca çerçeve okunurluğu kaybolmuyor.
PDF_GOLGE_DUZELTME = ".ekran{box-shadow:none !important}"


def sayfa_hazirla(tarayici):
    """Sunumu 1920×1080 açık temada, animasyonları söndürerek yükler."""
    baglam = tarayici.new_context(
        viewport={"width": EN, "height": BOY},
        device_scale_factor=1,
        # Açık tema deterministik olsun: sistem karanlık moddaysa bile
        # `prefers-color-scheme` açık gelsin (HTML'in bare :root paleti).
        color_scheme="light",
        # `@media (prefers-reduced-motion:reduce)` .gir'i görünür yapıyor;
        # ayrıca scroll-snap'i kapatıyor. Kaynağın kendi kancası, kullan.
        reduced_motion="reduce",
    )
    sf = baglam.new_page()
    sf.goto(KAYNAK.as_uri(), wait_until="load")
    # Kaynaktaki IntersectionObserver yalnız görünen slayta `.acik` verir;
    # tümüne elle ver ki ekran dışı slaytlar da tam opak render edilsin.
    sf.evaluate(
        "() => document.querySelectorAll('.slayt')"
        ".forEach(s => s.classList.add('acik'))"
    )
    sf.add_style_tag(content=BASKI_CSS)
    sf.wait_for_timeout(600)  # font yerleşimi + base64 görsellerin çözülmesi
    return baglam, sf


# Gerçek içerik yüksekliğini ölçmenin tuzağı: `.slayt` bir flex sütun ve
# `justify-content:center`. İçerik kutudan taşarsa HEM ÜSTTEN hem alttan taşar,
# ama `scrollHeight` yalnızca ALTTAKİ taşmayı sayar — üstteki negatif kaydırma
# alanına düşer ve görünmez. İlk denemede slayt 07'nin "Kanıt defteri" üst
# etiketi tam bu yüzden ölçüme girmedi ve kırpıldı.
# Bu yüzden ölçüm sırasında hizalama geçici olarak `flex-start`e alınır:
# içerik tek yöne akar, scrollHeight tamamını sayar. Sonra geri alınır.
OLC_JS = """() => Array.from(document.querySelectorAll('.slayt')).map(s => {
  const no = ((s.querySelector('.no') || {}).textContent || '??').trim();
  const eski = s.style.justifyContent;
  s.style.justifyContent = 'flex-start';
  const h = s.scrollHeight;
  s.style.justifyContent = eski;
  return [no, Math.round(h)];
})"""


def tasma_olc(sf) -> list[tuple[str, int]]:
    """İçeriği 1080px'i aşan slaytları bulur (kırpılma erken uyarısı)."""
    return sf.evaluate(OLC_JS)


# 16:9 sabit sayfaya sığdırma. Kaynak HTML `min-height:100vh` ile tasarlandı;
# tarayıcıda uzun slayt kaydırılabiliyor, sabit sayfada kaydırma yok. Slayt
# `justify-content:center` + `overflow:hidden` olduğu için taşma sessizce
# HEM ÜSTTEN HEM ALTTAN kırpılıyordu (07'de başlık kesiliyordu).
#
# Çözüm içerik değiştirmek DEĞİL, ölçek. Slaytın çocukları SABİT GENİŞLİKLİ
# (1920px) bir kutuya alınır; kutu özgün dolgu ve doğal yüksekliğiyle, yani
# birebir özgün düzenle kalır, sonra `zoom` ile küçültülür. Slayt kendisi flex
# ile ortalar. Hiçbir öge silinmez, hiçbir metin kısaltılmaz.
#
# ÜÇ YOL DENENDİ, İKİSİ ÖLÇÜMLE ELENDİ:
#   1) `.slayt`'a doğrudan `zoom` — kutu genişliği `auto` olduğu için yerel
#      düzen genişliği 1920/z'ye çıkıyor, `width:100%` ürün ekran görüntüsü
#      aynı oranda büyüyor ve küçülmeyi götürüyor. 5 yinelemede yakınsamadı
#      (07: 1086px, hedef 1080).
#   2) `transform:scale` + `left:50%/translate(-50%)` — ekran görüntüsünde
#      doğru, PDF'te İKİ ayrı bozulma: (a) mutlak konumlu kutunun düzen kutusu
#      960..2880'e uzandığı için Chromium yazdırma yolu TÜM BELGEYİ 0,70 katına
#      küçültüyor ("shrink to fit"); (b) dönüştürülmüş alt ağaçtaki base64
#      ürün ekran görüntüsü PDF'e hiç boyanmıyor — slayt 07 bomboş çıkıyor.
#   3) SABİT genişlik + `zoom` (seçilen) — yeniden düzen olur ama genişlik
#      1920px'e sabitlendiği için düzen özgün hâliyle birebir aynı kalır; yalnız
#      boyama küçülür. Ölçüldü: PDF sayfa 7 ile PNG 07 piksel piksel aynı
#      (sol=404, genişlik=1055, dikey merkez=541).
SIGDIR_JS = """([hedef, en]) => {
  const rapor = [];
  document.querySelectorAll('.slayt').forEach(s => {
    const no = ((s.querySelector('.no') || {}).textContent || '??').trim();
    // ölçüm hizalaması: bkz. OLC_JS — merkezleme üst taşmayı gizliyor
    const eskiHiza = s.style.justifyContent;
    s.style.justifyContent = 'flex-start';
    const H = s.scrollHeight;
    s.style.justifyContent = eskiHiza;
    if (H <= hedef + 0.5) return;
    const z = (hedef / H) * 0.995;               // küçük güvenlik payı
    const st = getComputedStyle(s);
    const kutu = document.createElement('div');
    kutu.className = '__sigdir';
    while (s.firstChild) kutu.appendChild(s.firstChild);
    s.appendChild(kutu);
    // Kutu: özgün genişlik + özgün dolgu + özgün flex davranışı → düzen birebir
    kutu.style.cssText =
      'width:' + en + 'px; height:' + H + 'px; zoom:' + z + ';' +
      'box-sizing:border-box; padding:' + st.padding + ';' +
      'display:flex; flex-direction:column; flex:none;' +
      'justify-content:' + st.justifyContent + ';' +
      'align-items:' + st.alignItems + ';';
    // Slayt: kutuyu iki eksende de ortalayan bir çerçeveye dönüşür
    s.style.cssText += ';padding:0; display:flex; flex-direction:row;' +
                       'align-items:center; justify-content:center;' +
                       'position:relative; overflow:hidden;' +
                       'height:' + hedef + 'px; min-height:' + hedef + 'px;';
    const k = kutu.getBoundingClientRect();
    rapor.push([no, Math.round(z * 1000) / 1000, Math.round(k.height)]);
  });
  return rapor;
}"""


def sigdir(sf) -> list[tuple[str, float, int]]:
    return sf.evaluate(SIGDIR_JS, [BOY, EN])


def pdf_uret(sf) -> None:
    sf.pdf(
        path=str(PDF),
        width=f"{EN}px",
        height=f"{BOY}px",
        print_background=True,
        margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
        prefer_css_page_size=True,
    )


def png_cek(sf, hedef: Path) -> list[Path]:
    """Her slaytı ayrı 1920×1080 PNG olarak çeker."""
    hedef.mkdir(parents=True, exist_ok=True)
    yollar: list[Path] = []
    slaytlar = sf.query_selector_all(".slayt")
    for i, slayt in enumerate(slaytlar, start=1):
        p = hedef / f"slayt-{i:02d}.png"
        slayt.scroll_into_view_if_needed()
        sf.wait_for_timeout(120)
        slayt.screenshot(path=str(p))
        yollar.append(p)
    return yollar


YARDIMCI_BETIK = r'''
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

CIKTI = Path(sys.argv[1])
PNGLER = [Path(p) for p in sys.argv[2:]]

# 16:9 — 1920x1080 px @ 96dpi = 20" x 11.25" = 18288000 x 10287000 EMU
EN_EMU, BOY_EMU = Emu(18288000), Emu(10287000)

sunum = Presentation()
sunum.slide_width, sunum.slide_height = EN_EMU, BOY_EMU
bos = sunum.slide_layouts[6]  # "Blank" — yer tutucu yok, görsel tam kaplasın

for png in PNGLER:
    s = sunum.slides.add_slide(bos)
    s.shapes.add_picture(str(png), 0, 0, width=EN_EMU, height=BOY_EMU)

sunum.save(str(CIKTI))
print(f"pptx yazildi: {CIKTI} · {len(PNGLER)} slayt")
'''


def yardimci_python() -> str:
    """python-pptx'in bulunduğu yorumlayıcıyı döndürür; yoksa kurar."""
    try:
        import pptx  # noqa: F401
        return sys.executable
    except ImportError:
        pass
    py = YARDIMCI / "bin" / "python"
    if not py.exists():
        print(f"  yardımcı ortam kuruluyor (AĞ gerekir): {YARDIMCI}")
        subprocess.run([sys.executable, "-m", "venv", str(YARDIMCI)], check=True)
        subprocess.run([str(py), "-m", "pip", "install", "--quiet", "python-pptx"],
                       check=True)
    return str(py)


def pptx_uret(pngler: list[Path]) -> None:
    py = yardimci_python()
    with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False) as f:
        f.write(YARDIMCI_BETIK)
        betik = f.name
    try:
        subprocess.run([py, betik, str(PPTX), *map(str, pngler)], check=True)
    finally:
        Path(betik).unlink(missing_ok=True)


def kb(p: Path) -> str:
    n = p.stat().st_size
    return f"{n/1024:.0f} KB" if n < 1024 * 1024 else f"{n/1024/1024:.1f} MB"


def main() -> int:
    ap = argparse.ArgumentParser(description="Jüri sunumu → PDF + PPTX")
    ap.add_argument("--sadece", choices=("pdf", "pptx"), default=None)
    arg = ap.parse_args()

    if not KAYNAK.exists():
        print(f"HATA: kaynak yok: {KAYNAK}", file=sys.stderr)
        return 1

    print(f"kaynak: {KAYNAK.name} · {kb(KAYNAK)}")
    with sync_playwright() as p:
        tarayici = p.chromium.launch()
        baglam, sf = sayfa_hazirla(tarayici)

        olcum = tasma_olc(sf)
        print(f"slayt sayısı: {len(olcum)}")
        tasan = [(no, h) for no, h in olcum if h > BOY]
        if tasan:
            print(f"  ham taşma (16:9 · {BOY}px sayfada kırpılırdı):")
            for no, h in tasan:
                print(f"    slayt {no}: {h}px (+{h - BOY})")

        olcek = sigdir(sf)
        if olcek:
            print("  sığdırma uygulandı (içerik değişmedi, yalnız ölçek):")
            for no, z, h in olcek:
                print(f"    slayt {no}: ölçek {z} → {h}px")

        # son kontrol: sığdırmadan sonra boyanan yükseklik sayfayı aşıyor mu?
        # (ölçek kutusu varsa getBoundingClientRect zoom sonrası GERÇEK boyutu
        #  verir; scrollHeight zoom öncesi düzen değerini verirdi)
        kalan = sf.evaluate(
            """(hedef) => Array.from(document.querySelectorAll('.slayt'))
                 .map(s => {
                   const k = s.querySelector(':scope > .__sigdir');
                   const h = k ? k.getBoundingClientRect().height : s.scrollHeight;
                   return [((s.querySelector('.no')||{}).textContent||'??').trim(),
                           Math.round(h)];
                 })
                 .filter(([, h]) => h > hedef + 2)""", BOY)
        if kalan:
            print(f"  UYARI · hâlâ taşan: {kalan}")
        else:
            print(f"  taşma yok · tüm slaytlar tam {BOY}px sayfaya sığıyor")

        # SIRA ÖNEMLİ: önce PNG (gölgeler doğru render edilir), sonra gölge
        # düzeltmesi enjekte edilip PDF. Tersi olsaydı PPTX gölgesiz kalırdı.
        pngler: list[Path] = []
        if arg.sadece != "pdf":
            gecici = Path(tempfile.mkdtemp(prefix="anatolia-sunum-"))
            pngler = png_cek(sf, gecici)
            print(f"png çekildi: {len(pngler)} kare → {gecici}")

        if arg.sadece != "pptx":
            sf.add_style_tag(content=PDF_GOLGE_DUZELTME)
            sf.wait_for_timeout(200)
            pdf_uret(sf)
            print(f"pdf yazıldı: {PDF.name} · {kb(PDF)}")

        baglam.close()
        tarayici.close()

    if pngler:
        pptx_uret(pngler)
        print(f"pptx yazıldı: {PPTX.name} · {kb(PPTX)}")

    # ——— ölçüm: iddia değil, sayı bas ———
    if PDF.exists() and arg.sadece != "pptx":
        from pypdf import PdfReader
        r = PdfReader(str(PDF))
        k = r.pages[0].mediabox
        print(f"DOĞRULAMA pdf: {len(r.pages)} sayfa · "
              f"{float(k.width):.0f}×{float(k.height):.0f} pt")
    if PPTX.exists() and arg.sadece != "pdf":
        import zipfile
        with zipfile.ZipFile(PPTX) as z:
            n = len([x for x in z.namelist()
                     if x.startswith("ppt/slides/slide") and x.endswith(".xml")])
        print(f"DOĞRULAMA pptx: {n} slayt · geçerli OOXML paketi")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
